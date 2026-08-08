from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("Anki_Garden_Development_Status_and_UI_Assessment.docx")
INK = "243229"
GREEN = "356B4A"
PALE_GREEN = "E7F0E9"
GOLD = "9A6A18"
PALE_GOLD = "F7EEDB"
RED = "9A3D3D"
PALE_RED = "F7E6E6"
GRAY = "66706A"
LIGHT = "F3F5F3"
WHITE = "FFFFFF"


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc.get_or_add_tcPr()
    tc_mar = tc.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        for i, cell in enumerate(row.cells):
            set_cell_width(cell, widths[i])
            margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    header_pr = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header_pr.append(repeat)


def keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


def add_bullet(doc: Document, text: str, *, bold_lead: str | None = None) -> None:
    p = doc.add_paragraph(style="List Bullet")
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)


def add_status_table(doc: Document, rows: list[tuple[str, str, str, str]], widths=None) -> None:
    widths = widths or [1850, 1450, 3300, 2760]
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = ["Feature area", "Status", "Current state", "Evidence / remaining gap"]
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade(cell, GREEN)
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        run.font.color.rgb = rgb(WHITE)
    for area, status, state, evidence in rows:
        cells = table.add_row().cells
        values = [area, status, state, evidence]
        for i, value in enumerate(values):
            cells[i].text = value
            if i == 1:
                cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                fill = PALE_GREEN if status == "Complete" else PALE_GOLD if status == "Partial" else PALE_RED
                color = GREEN if status == "Complete" else GOLD if status == "Partial" else RED
                shade(cells[i], fill)
                for run in cells[i].paragraphs[0].runs:
                    run.bold = True
                    run.font.size = Pt(9.5)
                    run.font.color.rgb = rgb(color)
    set_table_geometry(table, widths)
    doc.add_paragraph()


def add_rating_table(doc: Document, rows: list[tuple[str, str, str, str]]) -> None:
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for i, text in enumerate(["UI surface", "Rating", "What works now", "Main opportunity"]):
        shade(table.rows[0].cells[i], GREEN)
        run = table.rows[0].cells[i].paragraphs[0].add_run(text)
        run.bold = True
        run.font.color.rgb = rgb(WHITE)
    for surface, rating, works, opportunity in rows:
        cells = table.add_row().cells
        for i, value in enumerate([surface, rating, works, opportunity]):
            cells[i].text = value
        cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        shade(cells[1], PALE_GREEN if float(rating.split("/")[0]) >= 8.5 else PALE_GOLD)
        cells[1].paragraphs[0].runs[0].bold = True
    set_table_geometry(table, [1800, 1000, 3500, 3060])


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)
section.header_distance = Inches(0.35)
section.footer_distance = Inches(0.35)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(10.5)
normal.font.color.rgb = rgb(INK)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.1
for name, size, before, after in (("Title", 24, 0, 5), ("Subtitle", 12, 0, 12), ("Heading 1", 16, 15, 7), ("Heading 2", 13, 11, 5), ("Heading 3", 11, 8, 4)):
    style = styles[name]
    style.font.name = "Aptos Display" if name != "Subtitle" else "Aptos"
    style.font.size = Pt(size)
    style.font.color.rgb = rgb(GREEN if name != "Subtitle" else GRAY)
    style.font.bold = name != "Subtitle"
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = name.startswith("Heading") or name == "Title"
for name in ("List Bullet", "List Number"):
    style = styles[name]
    style.font.name = "Aptos"
    style.font.size = Pt(10.5)
    style.paragraph_format.left_indent = Inches(0.38)
    style.paragraph_format.first_line_indent = Inches(-0.19)
    style.paragraph_format.space_after = Pt(4)
    style.paragraph_format.line_spacing = 1.15

header = section.header.paragraphs[0]
header.text = "ANKI GARDEN  |  DEVELOPMENT STATUS"
header.style = styles["Normal"]
header.runs[0].font.size = Pt(8)
header.runs[0].font.bold = True
header.runs[0].font.color.rgb = rgb(GRAY)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run("Internal planning brief  •  July 12, 2026")
footer.runs[0].font.size = Pt(8)
footer.runs[0].font.color.rgb = rgb(GRAY)

p = doc.add_paragraph(style="Title")
p.add_run("Anki Garden\nDevelopment Status & UI Assessment")
p = doc.add_paragraph(style="Subtitle")
p.add_run("A current-state brief for choosing the next development step")

table = doc.add_table(rows=1, cols=3)
table.style = "Table Grid"
for i, (label, value) in enumerate((("Overall product", "Release-candidate core"), ("Automated baseline", "Full suite passing"), ("Recommended next move", "Focused polish + beta feedback"))):
    cell = table.rows[0].cells[i]
    shade(cell, PALE_GREEN if i != 2 else PALE_GOLD)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(label.upper() + "\n")
    r.bold = True
    r.font.size = Pt(8)
    r.font.color.rgb = rgb(GRAY)
    r = p.add_run(value)
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = rgb(GREEN if i != 2 else GOLD)
set_table_geometry(table, [3120, 3120, 3120])

doc.add_heading("Executive summary", level=1)
p = doc.add_paragraph()
p.add_run("Bottom line. ").bold = True
p.add_run("Anki Garden’s focused 2.1 experience is substantially complete. The review-to-growth loop, persistent local state, home surfaces, dashboard, plant interaction, settings, Plant Stories, achievements, accessibility behavior, asset fallbacks, and packaging have automated and isolated-Anki evidence. The product is no longer waiting on a foundational feature; it is ready for a controlled beta/release-readiness phase.")
p = doc.add_paragraph()
p.add_run("The largest remaining gaps are deliberate rather than accidental: ").bold = True
p.add_run("real-user usability evidence is still thin; long-term engagement and progression balance are not yet proven; visual motion remains limited by flattened artwork; and deferred systems such as deck mapping, shop/currency, social/cloud, focus timers, and exam mode are intentionally outside the focused product.")

doc.add_heading("Status at a glance", level=1)
add_status_table(doc, [
    ("Core study loop", "Complete", "Review answers create visible growth, daily progress, streak/vitality changes, quests, and milestone progression.", "Automated engine/reviewer coverage plus packaged live-review acceptance."),
    ("Garden persistence", "Complete", "Progress is local-first in user_files; settings use Anki config; restart, migration, repair, and rollback paths exist.", "v8 reset/migration, restart persistence, malformed-state recovery, and atomic-save evidence."),
    ("Home presence", "Complete", "Deck Browser and Overview show a compact preview with progress, weather, vitality, plants, Open Garden, Refresh/Retry behavior.", "Cross-surface refresh, fallback, accessibility, and isolated display acceptance."),
    ("Interactive dashboard", "Complete", "Responsive painted garden with selection, nurture, move/swap, cancel, one-level undo, milestone claims, quests, and achievements.", "Mouse/keyboard packaged QA and interaction regression tests."),
    ("Plant identity & stories", "Complete", "Generated/editable names and private milestone memories tied to study behavior.", "Rename, story timeline, accessibility, persistence, and packaged acceptance."),
    ("Settings & customization", "Complete", "Theme, motion, quality, daily goal, home visibility, preview, defaults, transactional save, and troubleshooting.", "Wide/narrow layout, rollback, restart, and config-validation evidence."),
    ("Accessibility foundation", "Complete", "Keyboard paths, focus visibility, semantic progress/status, reduced motion, labeled actions, and escape/tab behavior.", "Automated contracts and isolated accessibility-tree inspection."),
    ("Release engineering", "Complete", "Repeatable tests, asset audit, compilation, package build, ZIP/parity checks, and disposable-profile acceptance process.", "The full current suite passes. Exact counts and package hashes belong in the dated release ledger."),
    ("Real-user validation", "Partial", "Engineering QA is strong, but broad learner feedback and task-based usability evidence are not recorded.", "Run a small beta with first-run, daily-use, and return-after-a-week tasks."),
    ("Long-term progression tuning", "Partial", "Milestones and rewards work, but pacing, retention, quest difficulty, and vitality feel are not empirically tuned.", "Instrument privacy-safe local diagnostics or conduct diary testing before adding more systems."),
    ("Advanced foliage motion", "Incomplete", "Flattened plant art cannot safely sway without moving pots/soil/stems.", "Requires layered fixed-base and foliage assets for every species/stage."),
    ("Peripheral game systems", "Incomplete", "Shop/currency, events, mastery, cloud/social, deck mapping, focus timer, and exam mode are absent by design.", "Do not restart these until the focused core proves demand."),
])

doc.add_heading("Major features — complete", level=1)
complete_sections = [
    ("1. Review-driven growth and progression", ["Every supported review answer is counted once and contributes growth.", "The nurtured plant receives 80%; the remainder is deterministically shared without losing points.", "Daily goals, streak, garden vitality, quests, achievements, plant stages, and milestone plant choices form one coherent loop.", "Same-day progress survives restart; day rollover and synced-review catch-up avoid duplication."]),
    ("2. Calm, local-first garden experience", ["Progress remains in the add-on’s preserved user_files area; no cloud, social, card-content, or deck-name history is required.", "Bundled local art and fallbacks avoid remote-image dependencies.", "Malformed or older state is repaired, backed up, migrated/reset according to the v8 pre-release contract, or safely replaced."]),
    ("3. Home card and full dashboard", ["Deck Browser and Overview expose the same compact garden summary and one clear Open Garden action.", "The dashboard presents the scene, daily progress, quests, achievements, reward choices, feedback, and settings without a reviewer-native control.", "Mutations refresh already-open home surfaces so visible values stay consistent."]),
    ("4. Plant care, placement, and attachment", ["Plants can be selected, nurtured, moved into empty slots, swapped, cancelled, and undone once.", "Mouse and keyboard routes are available; save failure rolls the layout back.", "Plant names and milestone stories add identity without storing study content."]),
    ("5. Responsive settings and operational support", ["Settings are staged until Save and remain unchanged if persistence fails.", "The layout stacks at narrow widths; controls expose values and accessible names.", "Troubleshooting provides readable diagnostics and report-copy support."]),
]
for heading, bullets in complete_sections:
    doc.add_heading(heading, level=2)
    for item in bullets:
        add_bullet(doc, item)

doc.add_heading("Major features — partially complete", level=1)
add_status_table(doc, [
    ("Onboarding", "Partial", "Versioned first-review and Nurture guidance, transactional dismissal, legacy migration, accessibility, and restart persistence passed exact-package QA.", "Engineering acceptance is complete; the remaining gap is observation with new learners."),
    ("Progression balance", "Partial", "Mechanics and numeric progress are implemented and internally consistent.", "Thresholds, reward cadence, quest mix, and vitality decay need behavioral evidence over weeks, not only seeded QA."),
    ("Visual richness", "Partial", "A cohesive storybook-gouache catalog, weather overlays, grounded plants, themes, fallbacks, and reduced motion are present.", "Most plant assets are flattened; richer wind/foliage motion needs a dedicated layered-art pipeline."),
    ("Release confidence", "Partial", "Automated, packaged, migration, restart, mouse, keyboard, and accessibility acceptance are unusually strong.", "A fresh beta cohort and multi-platform/theme observation would reduce the remaining product-risk gap."),
])

doc.add_heading("Major features — incomplete or intentionally deferred", level=1)
p = doc.add_paragraph()
p.add_run("These are not release blockers for the focused core. ").bold = True
p.add_run("They should be treated as future product bets and only promoted after beta evidence identifies a real need.")
for item in [
    "Layered foliage wind animation with fixed pots, soil, and stem bases.",
    "Shop, spendable currency, weekly events, mastery, rare events, and passive rewards.",
    "Focus timer, exam mode, and deck-to-garden mapping.",
    "Cloud sync, social sharing, leaderboards, or remote artwork downloads.",
    "A broader plant inventory/collection-management surface beyond achievements, milestone choices, and the six-slot garden.",
]:
    add_bullet(doc, item)

doc.add_heading("UI state assessment", level=1)
p = doc.add_paragraph()
p.add_run("Rating scale: ").bold = True
p.add_run("10 = polished and release-ready; 8 = strong with targeted refinement; 6 = functional but visibly incomplete. Scores reflect current repository evidence and recent isolated Anki 26.05 acceptance—not a new external usability study.")
add_rating_table(doc, [
    ("Deck Browser / Overview card", "8.7/10", "Compact, coherent, accessible, resilient, and immediately actionable; values refresh after mutations.", "Test information density and visual hierarchy with first-time users; ensure dark/light theme polish across real collections."),
    ("Main garden scene", "8.8/10", "Strong visual identity, grounded assets, deterministic depth, clear selection, responsive rendering, and fallbacks.", "Improve affordance discovery without adding clutter; layered motion is the largest visual ceiling."),
    ("Plant action panel", "9.0/10", "Native buttons, keyboard focus, nurture/move/story/cancel actions, inline status, and undo avoid modal interruption.", "Observe whether action priority and move-state feedback are instantly understood by new users."),
    ("Progress / quests / achievements", "8.4/10", "Numeric progress, criteria, completion dates, rewards, and coherent terminology are present.", "Potential density and cognitive-load risk; validate which details learners actually use daily."),
    ("Plant Story dialog", "8.6/10", "Distinctive emotional layer with artwork, editable names, local dates, milestones, accessibility, and persistence.", "Empty/early-life states may need richer but restrained guidance; assess repeated-use value."),
    ("Settings", "8.8/10", "Responsive scroll layout, staged save, explicit preview-only controls, defaults, rollback, accessible values, troubleshooting.", "The boundary between saved appearance settings and demonstration controls deserves novice testing."),
    ("Keyboard & accessibility", "8.7/10", "Arrow/Enter/Space/Escape/Tab flows, visible focus, semantic progress/status, reduced motion, accessible names.", "Complete a screen-reader-led task pass and Windows keyboard QA before calling it exemplary."),
    ("Cross-surface consistency", "9.1/10", "Shared wording, values, assets, refresh behavior, and focused-core entry points reduce fragmentation.", "Keep this discipline as new features arrive; resist adding parallel controls."),
    ("Overall UI", "8.8/10", "Cohesive, calm, feature-complete, responsive, and technically well validated.", "Next gains come from learner observation and small hierarchy/onboarding refinements, not a redesign."),
])

doc.add_heading("Recommended next development step", level=1)
p = doc.add_paragraph()
p.add_run("Recommendation: run a focused beta-and-polish cycle before adding another major feature. ").bold = True
p.add_run("The codebase already supports the intended product. The highest-value unknown is whether everyday learners understand and enjoy it over repeated use.")
steps = [
    ("1", "Create a short beta protocol", "Recruit 5–10 everyday Anki users. Observe first launch, first review growth, opening the garden, nurturing/moving a plant, reading a story, changing a setting, and returning after several days."),
    ("2", "Capture friction by severity", "Separate blockers, confusion, visual polish, and feature requests. Prioritize repeated problems over one-off preferences."),
    ("3", "Ship one constrained polish pass", "Fix onboarding wording, hierarchy, focus/selection cues, density, and responsive edge cases supported by evidence. Preserve the focused entry-point model."),
    ("4", "Re-run exact-package acceptance", "Run the automated gates and disposable-profile live scenarios on the final artifact; record package hash and results."),
    ("5", "Choose only one next product bet", "After beta results, select either deeper plant attachment/progression or the layered foliage art project. Do not reopen multiple peripheral systems at once."),
]
table = doc.add_table(rows=1, cols=3)
table.style = "Table Grid"
for i, text in enumerate(("Step", "Action", "Outcome")):
    shade(table.rows[0].cells[i], GREEN)
    run = table.rows[0].cells[i].paragraphs[0].add_run(text)
    run.bold = True
    run.font.color.rgb = rgb(WHITE)
for num, action, outcome in steps:
    cells = table.add_row().cells
    cells[0].text, cells[1].text, cells[2].text = num, action, outcome
    cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cells[0].paragraphs[0].runs[0].bold = True
    shade(cells[0], PALE_GREEN)
set_table_geometry(table, [700, 2500, 6160])

doc.add_heading("Decision guardrails", level=2)
for item in [
    "Do not treat deferred systems as missing essentials; they were intentionally removed from the focused 2.1 contract.",
    "Do not perform another broad UI redesign without evidence of a recurring learner problem.",
    "Maintain local-first privacy, single-source state behavior, native controls, inline feedback, keyboard parity, and cross-surface consistency.",
    "For any release candidate, validate the exact packaged archive in a separately keyed, sync-disabled disposable Anki profile.",
]:
    add_bullet(doc, item)

doc.add_heading("Evidence basis and confidence", level=1)
p = doc.add_paragraph("This assessment synthesizes the current repository implementation and documentation, including the README, focused feature-evidence matrix, codebase audit, UI state and entry-point references, future-feature backlog, tests, and the latest package/live-acceptance records. The full current automated suite passes; exact counts and artifact hashes are recorded only in dated release evidence.")
p = doc.add_paragraph()
p.add_run("Confidence: high for implementation status; medium for product desirability and long-term usability. ").bold = True
p.add_run("The remaining uncertainty is primarily user evidence, not engineering completeness.")

doc.save(OUT)
print(OUT)
